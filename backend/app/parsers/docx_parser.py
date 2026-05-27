"""DOCX document parser using python-docx."""

import io
import logging
from typing import Any

from app.parsers.base import BaseParser, ParsedDocument

logger = logging.getLogger(__name__)


class DOCXParser(BaseParser):
    """Parser for Microsoft Word DOCX documents."""

    supported_extensions = [".docx"]
    supported_mimetypes = [
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ]

    async def parse(
        self,
        file_path: str | None = None,
        file_bytes: bytes | None = None,
        filename: str = "",
    ) -> ParsedDocument:
        """Parse a DOCX document and extract text, headings, tables, and images."""
        try:
            from docx import Document as DocxDocument
            from docx.opc.constants import RELATIONSHIP_TYPE as RT
        except ImportError:
            raise ImportError(
                "python-docx is required for DOCX parsing. "
                "Install it with: pip install python-docx"
            )

        try:
            if file_path:
                doc = DocxDocument(file_path)
            elif file_bytes:
                doc = DocxDocument(io.BytesIO(file_bytes))
            else:
                raise ValueError("Either file_path or file_bytes must be provided.")

            pages: list[dict[str, Any]] = []
            tables: list[dict[str, Any]] = []
            figures: list[dict[str, Any]] = []
            all_text_parts: list[str] = []
            current_page_text: list[str] = []
            current_page_num = 1

            # Extract metadata from core properties
            metadata: dict[str, Any] = {"format": "docx"}
            try:
                core_props = doc.core_properties
                metadata["title"] = core_props.title or ""
                metadata["author"] = core_props.author or ""
                metadata["created"] = str(core_props.created) if core_props.created else ""
                metadata["modified"] = str(core_props.modified) if core_props.modified else ""
                metadata["subject"] = core_props.subject or ""
                metadata["keywords"] = core_props.keywords or ""
            except Exception as e:
                logger.debug(f"Could not extract DOCX properties: {e}")

            # Process paragraphs
            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                style_name = paragraph.style.name if paragraph.style else ""

                # Check for page breaks (approximate — DOCX doesn't have explicit pages)
                has_page_break = False
                for run in paragraph.runs:
                    if run._element.xml and "w:br" in run._element.xml and 'w:type="page"' in run._element.xml:
                        has_page_break = True
                        break

                if has_page_break and current_page_text:
                    page_content = "\n".join(current_page_text)
                    pages.append({
                        "page_num": current_page_num,
                        "text": page_content,
                        "images": [],
                    })
                    all_text_parts.append(page_content)
                    current_page_text = []
                    current_page_num += 1

                if text:
                    # Annotate headings for structure preservation
                    if style_name.startswith("Heading"):
                        try:
                            level = int(style_name.replace("Heading", "").strip())
                            prefix = "#" * level + " "
                        except ValueError:
                            prefix = "# "
                        current_page_text.append(prefix + text)
                    else:
                        current_page_text.append(text)

            # Flush remaining page content
            if current_page_text:
                page_content = "\n".join(current_page_text)
                pages.append({
                    "page_num": current_page_num,
                    "text": page_content,
                    "images": [],
                })
                all_text_parts.append(page_content)

            metadata["page_count"] = current_page_num

            # Extract tables
            for table_idx, table in enumerate(doc.tables):
                rows: list[list[str]] = []
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    rows.append(cells)

                if rows:
                    tables.append({
                        "page_num": None,  # DOCX doesn't map tables to pages easily
                        "description": f"Table {table_idx + 1}",
                        "data": {
                            "rows": rows,
                            "num_rows": len(rows),
                            "num_cols": max(len(r) for r in rows) if rows else 0,
                        },
                    })

            # Extract images from document relationships
            try:
                for rel_id, rel in doc.part.rels.items():
                    if "image" in rel.reltype:
                        try:
                            image_part = rel.target_part
                            image_bytes = image_part.blob
                            content_type = image_part.content_type or ""
                            figures.append({
                                "page_num": None,
                                "description": f"Embedded image ({content_type})",
                                "image_data": image_bytes,
                                "content_type": content_type,
                            })
                        except Exception as e:
                            logger.debug(f"Failed to extract image {rel_id}: {e}")
            except Exception as e:
                logger.debug(f"Failed to iterate document relationships: {e}")

            raw_text = "\n\n".join(all_text_parts)

            return ParsedDocument(
                raw_text=raw_text,
                pages=pages,
                tables=tables,
                figures=figures,
                metadata=metadata,
            )

        except Exception as e:
            logger.error(f"Error parsing DOCX '{filename}': {e}")
            raise
